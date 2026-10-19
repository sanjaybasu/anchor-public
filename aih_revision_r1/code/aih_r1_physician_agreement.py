"""Physician review, Tasks 1 and 2 (Note S10): checklist ratings and agreement of the LLM judges with physicians.

Reads the returned rater workbooks (downloaded by aih_r1_probes_v2.fetch_rater_workbooks into
adjudication/returned/), the sample key (adjudication/key.json), and the judges' decisions on the same 32 texts.
The Task 2 texts were drawn from the main record files (records_c1c_<model>.jsonl, rep 0) before the framework
re-run, so the judges' decisions come from the judgment files of those texts: the checklist judgments
(judgments_c1c_<model>__*.jsonl, with the untruncated re-judgments for long Claude texts) for the bare
configuration, configuration A, and the rewrite, and the whole-text judgments (judgments_c1ccat_<model>__*.jsonl)
for the appended text. Agreement does not depend on which run produced a text, because the physicians and the
judges assessed the same text.

Physician consensus for an item: the decision of the majority of raters, with "Partly" counted as not addressed;
an item without a majority is unresolved and excluded from agreement. Sensitivity and specificity are for
"addressed", with the physician consensus as reference; κ is Cohen's κ between the judge decision and the
consensus. CIs come from 2,000 bootstrap resamples clustered on case (seed 20261015). Writes
results_aih_r1/physician_review.json.

Run: python aih_r1_physician_agreement.py [--fetch]
"""
import json
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

DATA = Path(__file__).resolve().parent / "results_aih_r1"
ADJ = DATA / "adjudication"
JUDGES = {"gpt-6-astra": "GPT-6 Astra", "gemini-3.1-pro-preview": "Gemini 3.1 Pro", "claude-opus-5-5": "Claude Opus 5.5"}
SEED, N_BOOT = 20261015, 2000


def load_jsonl(p):
    out = []
    if p.exists():
        for line in open(p):
            try:
                out.append(json.loads(line))
            except Exception:
                pass
    return out


def judge_vectors(mk: str, case_id: str, configuration: str) -> dict:
    """{judge: [0/1 per checklist item]} for the text that Task 2 showed (rep 0 of the main record file)."""
    if configuration == "append":
        files, arm = sorted(DATA.glob(f"judgments_c1ccat_{mk}__*.jsonl")), "append_concat"
    else:
        files, arm = sorted(DATA.glob(f"judgments_c1c_{mk}__*.jsonl")), configuration
        # untruncated re-judgments of texts longer than 16,000 characters replace the truncated ones
        files += sorted(DATA.glob(f"judgments_c1cfull_{mk}__*.jsonl")) + sorted(DATA.glob(f"judgments_c1cfull2_{mk}__*.jsonl"))
    out = {}
    for f in files:
        for j in load_jsonl(f):
            if not j.get("ok"):
                continue
            m, cid, a, rep = j["key"].split("|")
            if m == mk and cid == case_id and a in (arm, f"{arm}_full") and int(rep) == 0:
                out[j["judge"]] = j["parsed"]
    return out


def read_tabs(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    t1 = [r for r in wb["1 Checklists"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    t2 = [r for r in wb["2 Item decisions"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    return t1, t2


def cohen_kappa(pairs):
    n = len(pairs)
    if not n:
        return None
    po = sum(1 for a, b in pairs if a == b) / n
    pa, pb = sum(a for a, _ in pairs) / n, sum(b for _, b in pairs) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return (po - pe) / (1 - pe) if pe < 1 else None


def fleiss(rows):
    """rows: per item, counts per category (equal number of raters per item)."""
    rows = [r for r in rows if sum(r) > 1]
    if not rows:
        return None
    n = sum(rows[0])
    N = len(rows)
    p_j = [sum(r[j] for r in rows) / (N * n) for j in range(len(rows[0]))]
    P_i = [(sum(c * c for c in r) - n) / (n * (n - 1)) for r in rows]
    Pbar, Pe = sum(P_i) / N, sum(p * p for p in p_j)
    return (Pbar - Pe) / (1 - Pe) if Pe < 1 else None


def fleiss_boot(rows_by_case: dict) -> list:
    """Percentile 95% CI for Fleiss' kappa from 2,000 resamples of cases."""
    rng, cases_ = random.Random(SEED), list(rows_by_case)
    ks = []
    for _ in range(N_BOOT):
        rr = [r for c in (rng.choice(cases_) for _ in cases_) for r in rows_by_case[c]]
        k = fleiss(rr) if rr else None
        if k is not None:
            ks.append(k)
    ks.sort()
    return [ks[int(0.025 * len(ks))], ks[int(0.975 * len(ks)) - 1]] if ks else [None, None]


def metrics(pairs_by_case: dict) -> dict:
    """pairs: (judge_addressed, physician_addressed). Point estimates and case-clustered bootstrap CIs."""
    def est(pairs):
        tp = sum(1 for j, p in pairs if j and p)
        fn = sum(1 for j, p in pairs if not j and p)
        tn = sum(1 for j, p in pairs if not j and not p)
        fp = sum(1 for j, p in pairs if j and not p)
        return {"sens": tp / (tp + fn) if tp + fn else None, "spec": tn / (tn + fp) if tn + fp else None,
                "kappa": cohen_kappa(pairs), "agree": (tp + tn) / len(pairs) if pairs else None,
                "tp": tp, "fn": fn, "tn": tn, "fp": fp}
    allp = [x for v in pairs_by_case.values() for x in v]
    point = est(allp)
    rng, cases = random.Random(SEED), list(pairs_by_case)
    boots = defaultdict(list)
    for _ in range(N_BOOT):
        e = est([x for c in (rng.choice(cases) for _ in cases) for x in pairs_by_case[c]])
        for k in ("sens", "spec", "kappa", "agree"):
            if e[k] is not None:
                boots[k].append(e[k])
    out = {"n_items": len(allp), "n_cases": len(cases), **{k: point[k] for k in ("tp", "fn", "tn", "fp")},
           "n_pos": point["tp"] + point["fn"], "n_neg": point["tn"] + point["fp"]}
    for k in ("sens", "spec", "kappa", "agree"):
        b = sorted(boots[k])
        out[k] = {"est": point[k], "ci": [b[int(0.025 * len(b))], b[int(0.975 * len(b)) - 1]] if b else [None, None]}
    # a proportion of 0% or 100% has a Wilson interval, as in the main analysis
    for k, kk, nn in (("sens", point["tp"], out["n_pos"]), ("spec", point["tn"], out["n_neg"])):
        if nn and kk in (0, nn):
            out[k]["ci"], out[k]["ci_method"] = wilson(kk, nn), "Wilson"
    return out


def wilson(k: int, n: int, z: float = 1.959964) -> list:
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * ((p * (1 - p) / n + z * z / (4 * n * n)) ** 0.5) / d
    return [max(0.0, c - h), min(1.0, c + h)]


def main():
    if "--fetch" in sys.argv:
        from aih_r1_probes_v2 import fetch_rater_workbooks
        fetch_rater_workbooks()
    cases = json.load(open(DATA / "cohort1c_cases_checklists.json"))
    key = json.load(open(ADJ / "key.json"))
    code2case = {v: k for k, v in key["case_codes"].items()}
    raters = {p.stem.replace("_", " "): read_tabs(p) for p in sorted((ADJ / "returned").glob("Rater_*.xlsx"))}
    n_r = len(raters)
    maj = n_r // 2 + 1
    res = {"n_raters": n_r}

    # ── Task 1: checklist items ──
    t1 = defaultdict(dict)          # (case_id, item_no) -> rater -> (required, harm)
    missing = defaultdict(set)      # case_id -> raters who listed a missing action
    for r, (rows, _) in raters.items():
        for row in rows:
            cid = code2case.get(str(row[0]).strip())
            if cid is None:
                continue
            if row[2] in (None, "") and row[6]:
                missing[cid].add(r)
            elif row[2] not in (None, "") and row[4]:
                t1[(cid, int(row[2]))][r] = (str(row[4]).strip(), str(row[5] or "").strip())
    req = {"High": Counter(), "other": Counter()}
    harm_high = Counter()
    fl_rows = []
    fl_by_case = defaultdict(list)
    for (cid, i), d in t1.items():
        grade = cases[cid]["checklist"][i - 1].get("grade")
        g = "High" if grade == "High" else "other"
        votes = Counter(v[0] for v in d.values())
        if len(d) == n_r:
            fl_rows.append([votes.get("Required", 0), votes.get("Not required", 0), votes.get("Unsure", 0)])
            fl_by_case[cid].append(fl_rows[-1])
        top, k = votes.most_common(1)[0]
        req[g][top if k >= maj else "no majority"] += 1
        if g == "High":
            hv = Counter(v[1] for v in d.values() if v[1])
            if hv:
                h, hk = hv.most_common(1)[0]
                harm_high[h if hk >= maj else "no majority"] += 1
    res["task1"] = {"n_cases": len({c for c, _ in t1}), "n_items": len(t1),
                    "required_high": dict(req["High"]), "required_other": dict(req["other"]),
                    "harm_high": dict(harm_high), "fleiss_kappa_required": fleiss(fl_rows),
                    "fleiss_kappa_required_ci": fleiss_boot(fl_by_case),
                    "cases_with_missing_action_listed": {"k": len(missing), "n": len({c for c, _ in t1})}}
    n_high = sum(req["High"].values())
    if n_high:
        res["task1"]["high_rated_required"] = {"k": req["High"].get("Required", 0), "n": n_high,
                                              "p": req["High"].get("Required", 0) / n_high}

    # ── Task 2: item decisions ──
    texts = {t["text_id"]: t for t in key["texts"]}
    phys = defaultdict(dict)        # (text_id, item_no) -> rater -> addressed (1/0)
    for r, (_, rows) in raters.items():
        for row in rows:
            tid, ino, dec = str(row[0]).strip(), row[2], str(row[4] or "").strip().lower()
            if tid in texts and ino not in (None, "") and dec in ("yes", "partly", "no"):
                phys[(tid, int(ino))][r] = 1 if dec == "yes" else 0
    fl2 = [[sum(d.values()), len(d) - sum(d.values())] for d in phys.values() if len(d) == n_r]
    fl2_by_case = defaultdict(list)
    for (tid, ino), d in phys.items():
        if len(d) == n_r:
            fl2_by_case[texts[tid]["case_id"]].append([sum(d.values()), len(d) - sum(d.values())])
    groups = defaultdict(lambda: defaultdict(list))
    fp_kinds = Counter()
    unresolved = 0
    for (tid, ino), d in phys.items():
        t = texts[tid]
        votes = Counter(d.values())
        cons = 1 if votes.get(1, 0) >= maj else 0 if votes.get(0, 0) >= maj else None
        if cons is None:
            unresolved += 1
            continue
        jv = judge_vectors(t["model"], t["case_id"], t["configuration"])
        jdec = {j: v[ino - 1] for j, v in jv.items() if isinstance(v, list) and len(v) >= ino}
        if len(jdec) < 2:
            continue
        jmaj = 1 if sum(jdec.values()) >= 2 else 0
        if jmaj == 1 and cons == 0:     # judges credited an item that the physicians did not
            decs = [str(rr[4] or "").strip().lower() for rname, (_, rws) in raters.items() for rr in rws
                    if str(rr[0]).strip() == tid and rr[2] not in (None, "") and int(rr[2]) == ino]
            fp_kinds["any_no" if "no" in decs else "partly_only"] += 1
        cid = t["case_id"]
        groups["overall"][cid].append((jmaj, cons))
        groups[f"configuration:{t['configuration']}"][cid].append((jmaj, cons))
        groups[f"model:{t['model']}"][cid].append((jmaj, cons))
        for j, v in jdec.items():
            groups[f"judge:{j}"][cid].append((v, cons))
    res["task2"] = {"n_texts": len({tid for tid, _ in phys}), "n_items_rated": len(phys),
                    "unresolved": unresolved, "fleiss_kappa_physicians": fleiss(fl2),
                    "fleiss_kappa_physicians_ci": fleiss_boot(fl2_by_case),
                    "agreement": {g: metrics(v) for g, v in groups.items()},
                    "judge_credited_physician_not": dict(fp_kinds)}
    (DATA / "physician_review.json").write_text(json.dumps(res, indent=1))
    print(json.dumps({k: v for k, v in res.items() if k != "task2"}, indent=1))
    print("task2 overall:", json.dumps(res["task2"]["agreement"].get("overall"), indent=1)[:600])


if __name__ == "__main__":
    main()
